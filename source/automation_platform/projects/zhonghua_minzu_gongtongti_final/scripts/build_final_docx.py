from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


PROJECT_DIR = Path(__file__).resolve().parents[1]
PAPER_MD = PROJECT_DIR / "final_paper.md"
OUT_DOCX = PROJECT_DIR / "中华民族共同体概论期末论文.docx"
DESKTOP_DOCX = Path("C:/Users/SunYu/Desktop") / "中华民族共同体概论期末论文_可提交版.docx"


def set_run_font(run, font_name: str, size_pt: float, bold: bool | None = None) -> None:
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold


def set_paragraph_format(paragraph, line_spacing: float = 1.5, first_line: bool = False) -> None:
    paragraph.paragraph_format.line_spacing = line_spacing
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    if first_line:
        paragraph.paragraph_format.first_line_indent = Pt(24)


def add_centered_text(doc: Document, text: str, font: str, size: float, bold: bool = False) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(p)
    run = p.add_run(text)
    set_run_font(run, font, size, bold)


def set_cell_text(cell, text: str, font: str = "宋体", size: float = 12, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    set_paragraph_format(p)
    run = p.add_run(text)
    set_run_font(run, font, size, bold)


def set_table_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "8")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")


def add_cover(doc: Document) -> None:
    for _ in range(3):
        doc.add_paragraph()
    add_centered_text(doc, "2025-2026学年第二学期结课论文", "黑体", 18, True)
    for _ in range(4):
        doc.add_paragraph()
    add_centered_text(doc, "中华民族共同体概论", "黑体", 20, True)
    for _ in range(2):
        doc.add_paragraph()
    add_centered_text(doc, "中国式现代化进程中铸牢中华民族共同体意识的实践路径研究", "黑体", 16, True)
    for _ in range(4):
        doc.add_paragraph()

    rows = [
        ("课程名", "中华民族共同体概论"),
        ("论文题目", "中国式现代化进程中铸牢中华民族共同体意识的实践路径研究"),
        ("任课教师", "王文昊"),
        ("姓名", "________________"),
        ("学号", "________________"),
        ("班级", "________________"),
        ("年级", "2023级"),
        ("专业", "________________"),
        ("学院", "________________"),
        ("完成时间", "2026年06月10日"),
    ]
    table = doc.add_table(rows=len(rows), cols=3)
    table.autofit = True
    set_table_borders(table)
    for row, (label, value) in zip(table.rows, rows):
        set_cell_text(row.cells[0], label, "宋体", 12, True)
        set_cell_text(row.cells[1], "：", "宋体", 12)
        set_cell_text(row.cells[2], value, "宋体", 12)
    doc.add_page_break()


def parse_markdown() -> tuple[str, str, str, list[tuple[str, list[str]]], list[str]]:
    text = PAPER_MD.read_text(encoding="utf-8")
    lines = text.splitlines()
    title = lines[0].lstrip("# ").strip()
    abstract = ""
    keywords = ""
    sections: list[tuple[str, list[str]]] = []
    references: list[str] = []
    current_heading = ""
    current_paragraphs: list[str] = []
    mode = None
    buffer: list[str] = []

    def flush_buffer_to_mode() -> None:
        nonlocal abstract, keywords, buffer
        value = "\n".join(line.strip() for line in buffer if line.strip()).strip()
        if mode == "abstract":
            abstract = value
        elif mode == "keywords":
            keywords = value
        buffer = []

    def flush_section() -> None:
        nonlocal current_heading, current_paragraphs
        if current_heading:
            sections.append((current_heading, current_paragraphs))
        current_heading = ""
        current_paragraphs = []

    for line in lines[1:]:
        if line.startswith("## "):
            flush_buffer_to_mode()
            flush_section()
            heading = line[3:].strip()
            if heading == "摘要":
                mode = "abstract"
            elif heading == "关键词":
                mode = "keywords"
            elif heading == "参考文献":
                mode = "references"
            else:
                mode = "section"
                current_heading = heading
        elif mode in {"abstract", "keywords"}:
            buffer.append(line)
        elif mode == "section":
            if line.strip():
                current_paragraphs.append(line.strip())
        elif mode == "references":
            if line.strip():
                references.append(line.strip())

    flush_buffer_to_mode()
    flush_section()
    return title, abstract, keywords, sections, references


def add_toc(doc: Document, sections: list[tuple[str, list[str]]]) -> None:
    add_centered_text(doc, "目录", "黑体", 16, True)
    doc.add_paragraph()
    entries = ["摘要", "关键词"] + [heading for heading, _ in sections] + ["参考文献", "评分标准"]
    for entry in entries:
        p = doc.add_paragraph()
        set_paragraph_format(p)
        run = p.add_run(entry)
        set_run_font(run, "宋体", 12)
    doc.add_page_break()


def add_paper_body(doc: Document) -> None:
    title, abstract, keywords, sections, references = parse_markdown()
    add_toc(doc, sections)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(p)
    run = p.add_run(title)
    set_run_font(run, "黑体", 16, True)

    p = doc.add_paragraph()
    set_paragraph_format(p, first_line=False)
    label = p.add_run("摘要：")
    set_run_font(label, "黑体", 14, True)
    content = p.add_run(abstract)
    set_run_font(content, "楷体", 14)

    p = doc.add_paragraph()
    set_paragraph_format(p, first_line=False)
    label = p.add_run("关键词：")
    set_run_font(label, "黑体", 14, True)
    content = p.add_run(keywords)
    set_run_font(content, "楷体", 14)

    for heading, paragraphs in sections:
        p = doc.add_paragraph()
        set_paragraph_format(p)
        run = p.add_run(heading)
        set_run_font(run, "黑体", 12)
        for para in paragraphs:
            p = doc.add_paragraph()
            set_paragraph_format(p, first_line=True)
            run = p.add_run(para)
            set_run_font(run, "宋体", 12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(p)
    run = p.add_run("参考文献")
    set_run_font(run, "黑体", 16, True)
    for item in references:
        p = doc.add_paragraph()
        set_paragraph_format(p)
        run = p.add_run(item)
        set_run_font(run, "宋体", 10.5)


def scoring_rows() -> list[tuple[str, str, str]]:
    return [
        ("评价项目", "评价标准", "评价标准"),
        ("选题（10%）", "9-10分", "选题立意新颖，观点鲜明，有较强的现实意义或理论价值，充分体现选题要求，符合中华民族共同体课程主题。"),
        ("选题（10%）", "7-8分", "选题有一定新意，并有一定的现实意义或理论价值，对所论述的观点有较明确的认识，切合个人实际。"),
        ("选题（10%）", "5-6分", "选题没有新意，对本门课程现实意义不大，不太切合本门课程。"),
        ("选题（10%）", "4分以下", "选题与以往论文有较大雷同，对该选题缺乏明确认识，对现实意义不大，不切合个人实际。"),
        ("文献资料的引用（10%）", "9-10分", "所参考的文献资料充分；凡直接引用他人原文或间接借用他人观点之处，均明确地标注出处，且直接引用的文字不多于所撰写论文的30%，对所引文献资料的理解准确。"),
        ("文献资料的引用（10%）", "7-8分", "文献资料引用较合理；凡直接引用之处均明确标注出处；所引文字不超过全文的40%，对文献资料的理解正确。"),
        ("文献资料的引用（10%）", "5-6分", "文献资料不够充足，且引用目的不够明确，直接引用之处尚有明确标注，直接引用的文字多于40%，且对部分文献资料的理解存在一定的偏差。"),
        ("文献资料的引用（10%）", "4分以下", "文献资料明显不足，且某些文献资料与题目相关性不大，部分引用他人的文字未作说明，所直接引用的文字超过了全文的70%。"),
        ("论文规范性（10%）", "9-10分", "字数控制在4000-6000字之内，目录、摘要、关键词、正文、参考文献选录、字体以及编排方式均符合要求。"),
        ("论文规范性（10%）", "7-8分", "字数基本在所要求的范围之内，目录、摘要、关键词、正文符合要求，参考文献选录基本适当、字体、编排基本符合要求。"),
        ("论文规范性（10%）", "5-6分", "字数基本在所要求的范围之内，目录、摘要、关键词、正文基本符合要求，参考文献选录不够充分、字体、编排存在个别错误。"),
        ("论文规范性（10%）", "4分以下", "字数明显多于或少于所要求的范围，目录、摘要、关键词、正文等与要求不符，参考文献选录过少、字体、编排明显错误。"),
        ("论文创新性（20%）", "18-20分", "用新方法进行调查研究，采用的资料较新，研究结果有原创性。"),
        ("论文创新性（20%）", "14-17分", "研究视角有一定新意，能够合理运用文献资料提炼观点，论证基本清楚，结论合理。"),
        ("论文创新性（20%）", "10-13分", "研究方法及视角均无创新，但尚能从他人的观点中发现问题，结论有一定可靠性。"),
        ("论文创新性（20%）", "9分以下", "研究方法及研究视角均无创新之处，抄袭别人论文或之前已在刊物中采用，所得出的结论无明显价值。"),
        ("论证（20%）", "18-20分", "论证思路清楚，逻辑性强；显示了比较扎实的专业知识；能准确理解并合理利用所引文献，遵守学术规范，体现了作者良好的学风；研究所得结论可靠。"),
        ("论证（20%）", "14-17分", "论证思路比较清楚，有一定的逻辑性；对文献资料理解正确，引用比较合理，注解基本清楚；结论比较可靠。"),
        ("论证（20%）", "10-13分", "论证基本清楚，某些地方层次不够合理，个别地方引用不当，或表达不够清楚，但整体看来，尚可从所用材料中得到论文的结论。"),
        ("论证（20%）", "9分以下", "思路不清楚，论证不够严密，超过20%的篇幅表述不清楚，结论不够可靠。"),
        ("论文工作量（10%）", "9-10分", "所写论文需要花费较长的时间的调查研究，文献资料的阅读量比较大，需要较大投入。"),
        ("论文工作量（10%）", "7-8分", "所写论文需要较多的时间和精力的投入，文献资料的阅读有一定的难度，论文显示作者对所研究的问题进行了独立的思考。"),
        ("论文工作量（10%）", "5-6分", "所写论文无需花费长时间的调查研究，所参考的文献资料比较少，但论文显示作者尚能独立思考问题。"),
        ("论文工作量（10%）", "4分以下", "没有认真阅读文献，过多采用网络文献，且未清楚注明来源，有拼凑而成之嫌。"),
        ("语言（20%）", "18-20分", "语言符合论文语体特点，语句之间衔接自然，通顺连贯，遣词造句准确无误，无错别字。"),
        ("语言（20%）", "14-17分", "语言基本符合论文语体特点，语句比较连贯，表达基本清楚，有个别轻微语病，无错别字。"),
        ("语言（20%）", "10-13分", "语句基本通顺，有少量错别字和轻微语病。"),
        ("语言（20%）", "9分以下", "超过10%的篇幅采用非论述性语体，语句不够通畅，错别字在10个以上，语病在5处以上。"),
        ("总得分", "", ""),
        ("备注", "", "凡有以下情况者，合计得分不应超过60分：（1）引用超过80%，且未作注解说明出处；（2）整篇错别字累计超过15个；（3）语病超过10处。"),
    ]


def add_scoring_table(doc: Document) -> None:
    doc.add_page_break()
    add_centered_text(doc, "评分标准", "黑体", 16, True)
    rows = scoring_rows()
    table = doc.add_table(rows=len(rows), cols=3)
    table.autofit = True
    set_table_borders(table)
    for row, values in zip(table.rows, rows):
        for cell, value in zip(row.cells, values):
            set_cell_text(cell, value, "宋体", 9, row is table.rows[0])


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)

    add_cover(doc)
    add_paper_body(doc)
    add_scoring_table(doc)
    doc.save(OUT_DOCX)
    shutil.copy2(OUT_DOCX, DESKTOP_DOCX)
    print(OUT_DOCX)
    print(DESKTOP_DOCX)


if __name__ == "__main__":
    build()
